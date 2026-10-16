"""AU-ONLY partnership_stats_AU.xlsx (Partnerships + Sites + Notes), same layout as h_report.py. Offline."""
import os, csv, json
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
D = os.environ.get('OUT', '/tmp/au_out'); B = json.load(open(f'{D}/au_classify_meta.json'))['built']
def num(v):
    try: return float(v) if v != '' else v
    except Exception: return v
wb = Workbook(); first = True
for name, path in [('Partnerships', f'{D}/partnership_stats_AU.csv'), ('Sites', f'{D}/partnership_sites_AU.csv')]:
    ws = wb.active if first else wb.create_sheet(); ws.title = name; first = False
    with open(path, encoding='utf-8') as fh:
        for i, row in enumerate(csv.reader(fh)):
            ws.append(row if i == 0 else [num(c) if len(c) < 30 else c[:32000] for c in row])
    ws.freeze_panes = 'D2' if name == 'Partnerships' else 'C2'; ws.auto_filter.ref = ws.dimensions
    for c in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(c)].width = 14
ws = wb.create_sheet('Notes')
ws.append([f"AU-ONLY: Australian (AUD) sites only. Built offline {B} from existing data on disk (no scraping/network). Method: AU = country AU or currency AUD; blank-country sites only via exact merchant_id/host/sites_config match; groups re-cut to AU members, <2 AU sites dropped, all stats recomputed on AU members. 'Unknown' country sites are excluded (see partnership_sites_unknown_country.csv)."])
ws.append([''])
for line in open(f'{D}/partnership_summary_AU.md', encoding='utf-8'): ws.append([line.rstrip('\n')])
ws.column_dimensions['A'].width = 160
wb.save(f'{D}/partnership_stats_AU.xlsx'); print('xlsx ok')
