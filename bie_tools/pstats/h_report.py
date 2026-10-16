"""partnership_stats.xlsx (Partnerships + Sites + Notes sheets)."""
import csv, json
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
T = '/workspace/bie_tools'
def num(v):
    try: return float(v) if v != '' else v
    except Exception: return v
wb = Workbook(); first = True
for name, path in [('Partnerships', f'{T}/partnership_stats.csv'), ('Sites', f'{T}/partnership_sites.csv')]:
    ws = wb.active if first else wb.create_sheet(); ws.title = name; first = False
    with open(path, encoding='utf-8') as fh:
        for i, row in enumerate(csv.reader(fh)):
            ws.append(row if i == 0 else [num(c) if len(c) < 30 else c[:32000] for c in row])
    ws.freeze_panes = 'D2' if name == 'Partnerships' else 'C2'; ws.auto_filter.ref = ws.dimensions
    for c in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(c)].width = 14
ws = wb.create_sheet('Notes')
for line in open(f'{T}/partnership_summary.md', encoding='utf-8'): ws.append([line.rstrip('\n')])
ws.column_dimensions['A'].width = 160
wb.save(f'{T}/partnership_stats.xlsx'); print('xlsx ok')
